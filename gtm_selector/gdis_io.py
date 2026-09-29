"""Импорт данных ГДИС (гидродинамических исследований) из Excel и привязка к скважинам.

Формат файла ГДИС может отличаться от раза к разу (разные месторождения/
выгрузки), поэтому колонки ищутся гибко — по ключевым словам-синонимам,
содержащимся в нормализованном заголовке (см. ``normalize_header`` в
data_io.py), а не по точному названию. Часть колонок может отсутствовать
вообще — это нормально, не ошибка; единственная обязательная колонка —
номер скважины.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from .data_io import normalize_header
from .models import GdisRecord, Well, gdis_matches_well_object

# Ключевые слова-синонимы для гибкого поиска колонок: колонка считается
# найденной, если её нормализованный заголовок СОДЕРЖИТ хотя бы одно из слов.
_FIELD_KEYWORDS: dict[str, list[str]] = {
    "well_number": ["скв", "скважина"],
    "object_name": ["объект"],
    "horizon": ["горизонт"],
    "date": ["дата начала", "дата исслед", "дата"],
    "research_type": ["вид исслед", "вид работ", "тип исслед"],
    "interval": ["интервал перфор", "интервал"],
    "skin": ["скин"],
    "permeability": ["проницаем"],
    "quality": ["качество интерпретац", "качество"],
    "comment": ["коментар", "комментар", "примечан"],
}

_FIELD_LABELS: dict[str, str] = {
    "well_number": "Скважина",
    "object_name": "Объект",
    "horizon": "Горизонт",
    "date": "Дата",
    "research_type": "Вид исследования",
    "interval": "Интервал",
    "skin": "Скин",
    "permeability": "Проницаемость",
    "quality": "Качество интерпретации",
    "comment": "Комментарий",
}


def _find_column(headers_norm: list[str], keywords: list[str]) -> int | None:
    for idx, header in enumerate(headers_norm):
        if any(kw in header for kw in keywords):
            return idx
    return None


def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _str_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _float_or_none(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in ("", "-", "—"):
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def _date_or_empty(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return ""
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def _match_well(value, wells: list[Well]) -> Well | None:
    """Сопоставляет значение ячейки номера скважины со скважиной.

    Если значение — чистое число (в т.ч. записанное как 23.0 из ячейки
    Excel) — ищет скважину, чьё имя/номер оканчивается на это число после
    дефиса (аналогично сопоставлению LAS-файлов, см. las_io.match_las_to_well).
    Если значение — текст с буквами ("УС-10") — ищет полное совпадение имени
    или номера скважины (без учёта регистра/пробелов).
    """
    text = _str_cell(value)
    if not text:
        return None

    try:
        number = int(float(text.replace(",", ".")))
    except ValueError:
        number = None

    if number is not None:
        suffix = f"-{number}"
        for well in wells:
            if well.name.endswith(suffix) or well.id.endswith(suffix):
                return well
        return None

    text_lower = text.lower()
    for well in wells:
        if well.name.strip().lower() == text_lower or well.id.strip().lower() == text_lower:
            return well
    return None


def load_gdis_database(path: str | Path, wells: list[Well]) -> tuple[dict[str, list[GdisRecord]], list[str]]:
    """Читает базу ГДИС (Excel) и сопоставляет записи со скважинами.

    Ищет лист с «гдис» в названии (без учёта регистра); если не найден —
    берёт первый лист. Первая строка — заголовки, колонки ищутся по
    ключевым словам-синонимам (гибко, без точного совпадения названия).
    Обязательная колонка — номер скважины; остальные могут отсутствовать,
    это не ошибка.

    Возвращает (``{well_id: [GdisRecord, ...]}``, список предупреждений).
    Предупреждения включают: список ненайденных желательных колонок,
    строки с несопоставленным номером скважины и записи, относящиеся к
    другому объекту/горизонту скважины (см. ``gdis_matches_well_object``).
    """
    import openpyxl  # опциональная зависимость, нужна только для импорта Excel

    warnings: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet_name = next((s for s in wb.sheetnames if "гдис" in s.lower()), wb.sheetnames[0])
        sheet = wb[sheet_name]

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers_norm = [normalize_header(c) for c in header_row]

        col_index = {
            field: _find_column(headers_norm, keywords)
            for field, keywords in _FIELD_KEYWORDS.items()
        }

        if col_index["well_number"] is None:
            actual_headers = ", ".join(str(c) for c in header_row if c is not None)
            raise ValueError(
                "На листе ГДИС не найдена колонка с номером скважины. "
                f"Заголовки файла: {actual_headers}"
            )

        for field in _FIELD_KEYWORDS:
            if field != "well_number" and col_index[field] is None:
                warnings.append(
                    f"В файле ГДИС не найдена колонка «{_FIELD_LABELS[field]}» — "
                    "это поле будет пустым для всех записей"
                )

        records_by_well: dict[str, list[GdisRecord]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            well_cell = _cell(row, col_index["well_number"])
            if well_cell is None or _str_cell(well_cell) == "":
                continue

            well = _match_well(well_cell, wells)
            if well is None:
                warnings.append(
                    f"ГДИС: не найдена скважина для значения «{well_cell}» — строка пропущена"
                )
                continue

            record = GdisRecord(
                well_id=well.id,
                object_name=_str_cell(_cell(row, col_index["object_name"])),
                horizon=_str_cell(_cell(row, col_index["horizon"])),
                interval=_str_cell(_cell(row, col_index["interval"])),
                date=_date_or_empty(_cell(row, col_index["date"])),
                research_type=_str_cell(_cell(row, col_index["research_type"])),
                skin=_float_or_none(_cell(row, col_index["skin"])),
                permeability=_float_or_none(_cell(row, col_index["permeability"])),
                quality=_str_cell(_cell(row, col_index["quality"])),
                comment=_str_cell(_cell(row, col_index["comment"])),
            )

            if not gdis_matches_well_object(record, well):
                date_text = record.date or "дата неизвестна"
                warnings.append(
                    f"Скважина {well.id}: запись ГДИС от {date_text} относится к объекту "
                    f"{record.object_name or '—'}/{record.horizon or '—'} - отличается от "
                    f"текущего {well.formation or '—'}/{well.horizon or '—'}, возможен "
                    "перевод на другой интервал"
                )

            records_by_well.setdefault(well.id, []).append(record)
    finally:
        wb.close()

    return records_by_well, warnings
