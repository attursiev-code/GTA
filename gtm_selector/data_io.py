"""Импорт/экспорт скважин (CSV, Excel-база) и сохранение/загрузка проекта (JSON)."""

from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from pathlib import Path

from .models import Candidate, ProductionPoint, Well
from .params import Settings

_INTERVAL_RE = re.compile(r"(\d+(?:,\d+)?)\s*-\s*(\d+(?:,\d+)?)")


def parse_perforation_intervals(text: str) -> list[tuple[float, float]]:
    """Разбирает текст интервала перфорации на подынтервалы (кровля, подошва).

    Формат: подынтервалы разделены ';', внутри подынтервала — 'кровля-подошва',
    запятая внутри числа — десятичный разделитель (например,
    ``'1003-1005,3; 1025-1028'`` -> ``[(1003.0, 1005.3), (1025.0, 1028.0)]``).
    Нераспознанные фрагменты просто пропускаются, ошибку не вызывают.
    """
    if not text:
        return []
    intervals = []
    for top_str, bottom_str in _INTERVAL_RE.findall(text):
        top = float(top_str.replace(",", "."))
        bottom = float(bottom_str.replace(",", "."))
        intervals.append((top, bottom))
    return intervals

# Колонки листа "Добыча" в реальной базе (ищем по названию, не по номеру).
_PROD_COL_WELL = "Скважина"
_PROD_COL_PERF_INTERVAL = "Интервал перфорации"
_PROD_COL_HORIZON_OFM = "Горизонт по ОФМ"
_PROD_COL_HORIZON = "Горизонт"
_PROD_COL_OBJECT = "Объект"
_PROD_COL_DATE = "Дата"
_PROD_COL_QO = ("Дебит нефти т/сут", "Дебит нефти, т/сут")
_PROD_COL_QL = ("Дебит жидкости т/сут", "Дебит жидкости, т/сут")
_PROD_COL_WATERCUT = ("Обводненность %", "Обводненность, %", "Обводнённость %", "Обводнённость, %")

_COORD_SHEET = "Координаты"
_COORD_COL_WELL = "Скважина"
_COORD_COL_X = "X"
_COORD_COL_Y = "Y"
_COORD_OUTLIER_RATIO = 50.0  # во сколько раз порядок величины должен отличаться, чтобы считаться выбросом

WELL_FIELDS = [
    "id", "name", "formation", "status",
    "qo", "ql", "watercut",
    "p_res", "p_res_initial", "p_wf",
    "skin", "perm", "thickness_total", "thickness_perforated",
    "reserves_remaining", "depletion",
    "pump_capacity", "idle_days", "last_active_qo", "months_since_last_gtm",
    "notes",
]

_NUMERIC_FIELDS = {
    "qo", "ql", "watercut", "p_res", "p_res_initial", "p_wf", "skin", "perm",
    "thickness_total", "thickness_perforated", "reserves_remaining",
    "depletion", "pump_capacity", "last_active_qo",
}
_INT_FIELDS = {"idle_days", "months_since_last_gtm"}


def load_wells_csv(path: str | Path) -> list[Well]:
    wells: list[Well] = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            data = {}
            for key, value in row.items():
                if key is None:
                    continue
                key = key.strip()
                if value is None:
                    value = ""
                value = value.strip()
                if key in _NUMERIC_FIELDS:
                    data[key] = float(value) if value != "" else 0.0
                elif key in _INT_FIELDS:
                    data[key] = int(float(value)) if value != "" else 0
                else:
                    data[key] = value
            wells.append(Well.from_dict(data))
    return wells


def load_history_csv(path: str | Path) -> dict[str, list[ProductionPoint]]:
    """Читает историю добычи (well_id, date, qo, qw) и группирует по скважинам."""
    history: dict[str, list[ProductionPoint]] = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            well_id = (row.get("well_id") or "").strip()
            if not well_id:
                continue
            date = (row.get("date") or "").strip()
            qo_raw = (row.get("qo") or "").strip()
            qw_raw = (row.get("qw") or "").strip()
            qo = float(qo_raw) if qo_raw != "" else 0.0
            qw = float(qw_raw) if qw_raw != "" else 0.0
            history.setdefault(well_id, []).append(ProductionPoint(date=date, qo=qo, qw=qw))
    return history


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _cell_to_date_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text or None


def _header_index(header_row) -> dict[str, int]:
    return {str(cell).strip(): idx for idx, cell in enumerate(header_row) if cell is not None}


def _find_col(col_index: dict[str, int], *names: str) -> int | None:
    for name in names:
        if name in col_index:
            return col_index[name]
    return None


def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _load_coordinates_sheet(wb, well_ids: set[str]) -> tuple[dict[str, tuple[float, float]], list[str]]:
    """Читает лист «Координаты» (Скважина, X, Y), отфильтровывая явные выбросы
    по порядку величины (например метры/проекция среди координат в градусах)."""
    warnings: list[str] = []
    if _COORD_SHEET not in wb.sheetnames:
        return {}, warnings

    sheet = wb[_COORD_SHEET]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    col_index = _header_index(header_row)
    col_well = _find_col(col_index, _COORD_COL_WELL)
    col_x = _find_col(col_index, _COORD_COL_X)
    col_y = _find_col(col_index, _COORD_COL_Y)
    if col_well is None or col_x is None or col_y is None:
        warnings.append(
            "Лист «Координаты»: не найдены обязательные колонки (Скважина, X, Y) — "
            "координаты не загружены"
        )
        return {}, warnings

    raw: dict[str, tuple[float, float]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        well_id = _cell(row, col_well)
        if well_id is None or str(well_id).strip() == "":
            continue
        well_id = str(well_id).strip()
        x = _to_float(_cell(row, col_x))
        y = _to_float(_cell(row, col_y))
        raw[well_id] = (x, y)

    magnitudes = sorted(max(abs(x), abs(y)) for x, y in raw.values() if max(abs(x), abs(y)) > 0)
    median_mag = magnitudes[len(magnitudes) // 2] if magnitudes else 0.0

    coords: dict[str, tuple[float, float]] = {}
    for well_id, (x, y) in raw.items():
        mag = max(abs(x), abs(y))
        is_outlier = (
            median_mag > 0
            and mag > 0
            and (mag / median_mag > _COORD_OUTLIER_RATIO or median_mag / mag > _COORD_OUTLIER_RATIO)
        )
        if is_outlier:
            warnings.append(
                f"Скважина {well_id}: координаты ({x}, {y}) на порядки отличаются от "
                f"большинства остальных скважин (медиана {median_mag:.1f}) — "
                "исключена из словаря координат"
            )
            continue
        coords[well_id] = (x, y)

    for well_id in sorted(well_ids - set(raw.keys())):
        warnings.append(f"Скважина {well_id}: координаты отсутствуют на листе «Координаты»")

    return coords, warnings


def load_field_database(
    path: str | Path,
) -> tuple[list[Well], dict[str, tuple[float, float]], list[str]]:
    """Читает реальную базу промысловых данных (Excel, отдельный файл).

    Лист «Добыча» (или первый лист, если он называется иначе) группируется
    по скважине: строки сортируются по дате и превращаются в историю добычи
    (``ProductionPoint``), а текущие Qн/Qж/обводнённость и geo-поля скважины
    берутся из последней по дате строки. Лист «Координаты» — необязателен.

    Возвращает (список скважин, словарь {well_id: (x, y)}, список предупреждений).
    """
    import openpyxl  # опциональная зависимость, нужна только для импорта Excel-базы

    warnings: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = wb["Добыча"] if "Добыча" in wb.sheetnames else wb[wb.sheetnames[0]]

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        col_index = _header_index(header_row)
        col_well = _find_col(col_index, _PROD_COL_WELL)
        col_perf = _find_col(col_index, _PROD_COL_PERF_INTERVAL)
        col_horizon_ofm = _find_col(col_index, _PROD_COL_HORIZON_OFM)
        col_horizon = _find_col(col_index, _PROD_COL_HORIZON)
        col_object = _find_col(col_index, _PROD_COL_OBJECT)
        col_date = _find_col(col_index, _PROD_COL_DATE)
        col_qo = _find_col(col_index, *_PROD_COL_QO)
        col_ql = _find_col(col_index, *_PROD_COL_QL)
        col_watercut = _find_col(col_index, *_PROD_COL_WATERCUT)

        if col_well is None or col_date is None or col_qo is None:
            raise ValueError(
                "На листе «Добыча» не найдены обязательные колонки "
                "(Скважина, Дата, Дебит нефти т/сут)"
            )

        rows_by_well: dict[str, list[dict]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            well_id = _cell(row, col_well)
            if well_id is None or str(well_id).strip() == "":
                continue
            well_id = str(well_id).strip()

            date_str = _cell_to_date_str(_cell(row, col_date))
            if date_str is None:
                continue

            qo = _to_float(_cell(row, col_qo))
            ql = _to_float(_cell(row, col_ql))
            watercut = _to_float(_cell(row, col_watercut))

            perf_val = _cell(row, col_perf)
            horizon_val = _cell(row, col_horizon)
            horizon_ofm_val = _cell(row, col_horizon_ofm)
            object_val = _cell(row, col_object)

            rows_by_well.setdefault(well_id, []).append({
                "date": date_str,
                "qo": qo,
                "ql": ql,
                "watercut": watercut,
                "perforation_interval": str(perf_val).strip() if perf_val is not None else "",
                "horizon": str(horizon_val).strip() if horizon_val is not None else "",
                "horizon_ofm": str(horizon_ofm_val).strip() if horizon_ofm_val is not None else "",
                "object": str(object_val).strip() if object_val is not None else "",
            })

        wells: list[Well] = []
        for well_id, rows in rows_by_well.items():
            rows.sort(key=lambda r: r["date"])
            history = [
                ProductionPoint(date=r["date"], qo=r["qo"], qw=max(0.0, r["ql"] - r["qo"]))
                for r in rows
            ]
            last = rows[-1]
            formation = last["object"] or last["horizon"]
            horizon = last["horizon"] or last["horizon_ofm"]
            wells.append(Well(
                id=well_id,
                name=well_id,
                formation=formation,
                perforation_interval=last["perforation_interval"],
                horizon=horizon,
                qo=last["qo"],
                ql=last["ql"],
                watercut=last["watercut"],
                history=history,
            ))

        coords, coord_warnings = _load_coordinates_sheet(wb, {w.id for w in wells})
        warnings.extend(coord_warnings)
    finally:
        wb.close()

    return wells, coords, warnings


def save_wells_csv(wells: list[Well], path: str | Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=WELL_FIELDS)
        writer.writeheader()
        for well in wells:
            d = well.to_dict()
            writer.writerow({k: d.get(k, "") for k in WELL_FIELDS})


def save_recommendations_csv(recs: list[Candidate], path: str | Path) -> None:
    fieldnames = ["well_id", "well_name", "gtm_type", "gtm_label", "mechanism", "reasons"]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in recs:
            writer.writerow({
                "well_id": r.well_id,
                "well_name": r.well_name,
                "gtm_type": r.gtm_type.value,
                "gtm_label": r.gtm_label,
                "mechanism": r.mechanism,
                "reasons": "; ".join(r.reasons),
            })


def save_project(path: str | Path, wells: list[Well], settings: Settings) -> None:
    payload = {
        "wells": [w.to_dict() for w in wells],
        "settings": settings.to_dict(),
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def load_project(path: str | Path) -> tuple[list[Well], Settings]:
    with open(path, encoding="utf-8") as f:
        payload = json.load(f)
    wells = [Well.from_dict(d) for d in payload.get("wells", [])]
    settings = Settings.from_dict(payload.get("settings", {}))
    return wells, settings
